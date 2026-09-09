import hashlib
import mimetypes
from pathlib import Path

import pandas as pd
from pypdf import PdfReader
from PIL import Image
from PIL.ExifTags import TAGS, GPSTAGS
from mutagen import File as MutagenFile
from hachoir.parser import createParser
from hachoir.metadata import extractMetadata

def calculate_file_hash(file_path: str) -> str:
    sha256 = hashlib.sha256()

    with open(file_path, "rb") as file:
        while chunk := file.read(4096):
            sha256.update(chunk)

    return sha256.hexdigest()

def get_file_extension(filename: str) -> str:
    return Path(filename).suffix.lower()

IMAGE_EXTENSIONS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".webp",
    ".bmp",
    ".tiff"
}

AUDIO_EXTENSIONS = {
    ".mp3",
    ".wav",
    ".m4a",
    ".flac",
    ".ogg"
}

VIDEO_EXTENSIONS = {
    ".mp4",
    ".mov",
    ".avi",
    ".mkv",
    ".webm"
}

def _extract_gps(exif_data: dict):
    gps_info = exif_data.get("GPSInfo")

    if not gps_info:
        return None

    gps_data = {}

    for key, value in gps_info.items():
        tag_name = GPSTAGS.get(key, key)
        gps_data[tag_name] = str(value)

    return gps_data

def extract_image_metadata(file_path: str) -> dict:
    metadata = {
        "file_type": "image",
        "width": None,
        "height": None,
        "format": None,
        "mode": None,
        "exif": {}
    }

    try:
        with Image.open(file_path) as image:
            metadata["width"] = image.width
            metadata["height"] = image.height
            metadata["format"] = image.format
            metadata["mode"] = image.mode

            exif_data = image.getexif()

            if exif_data:
                exif = {}

                for tag_id, value in exif_data.items():
                    tag_name = TAGS.get(tag_id, tag_id)

                    if isinstance(value, bytes):
                        continue

                    exif[str(tag_name)] = str(value)

                gps = _extract_gps(exif_data)

                if gps:
                    exif["GPS"] = gps

                metadata["exif"] = exif

    except Exception as error:
        metadata["error"] = str(error)

    return metadata

def extract_basic_file_metadata(file_path: str, file_type: str) -> dict:
    path = Path(file_path)

    mime_type, _ = mimetypes.guess_type(file_path)

    return {
        "file_type": file_type,
        "size_bytes": path.stat().st_size if path.is_file() else None,
        "mime_type": mime_type
    }

def extract_audio_metadata(file_path: str) -> dict:
    path = Path(file_path)

    mime_type, _ = mimetypes.guess_type(file_path)

    metadata = {
        "file_type": "audio",
        "size_bytes": path.stat().st_size if path.is_file() else None,
        "mime_type": mime_type,
        "duration": None,
        "bitrate": None,
        "sample_rate": None,
        "channels": None,
        "format": path.suffix.lstrip(".").upper()
    }

    try:
        audio = MutagenFile(file_path)

        if audio is None:
            metadata["error"] = "Unsupported or unreadable audio file."
            return metadata

        info = audio.info

        metadata["duration"] = getattr(info, "length", None)
        metadata["bitrate"] = getattr(info, "bitrate", None)
        metadata["sample_rate"] = getattr(info, "sample_rate", None)
        metadata["channels"] = getattr(info, "channels", None)

    except Exception as error:
        metadata["error"] = str(error)

    return metadata

def extract_video_metadata(file_path: str) -> dict:
    path = Path(file_path)

    mime_type, _ = mimetypes.guess_type(file_path)

    metadata = {
        "file_type": "video",
        "size_bytes": path.stat().st_size if path.is_file() else None,
        "mime_type": mime_type,
        "duration": None,
        "width": None,
        "height": None,
        "codec": None,
        "fps": None,
        "bitrate": None,
        "format": path.suffix.lstrip(".").upper()
    }

    try:
        parser = createParser(str(file_path))

        if parser is None:
            metadata["error"] = "Unsupported or unreadable video file."
            return metadata

        hachoir_metadata = extractMetadata(parser)

        if hachoir_metadata is None:
            metadata["error"] = "Could not extract video metadata."
            return metadata

        if hachoir_metadata.has("duration"):
            metadata["duration"] = (
                hachoir_metadata.get("duration").total_seconds()
            )

        if hachoir_metadata.has("width"):
            metadata["width"] = hachoir_metadata.get("width")

        if hachoir_metadata.has("height"):
            metadata["height"] = hachoir_metadata.get("height")

        if hachoir_metadata.has("frame_rate"):
            metadata["fps"] = hachoir_metadata.get("frame_rate")

        if hachoir_metadata.has("bit_rate"):
            metadata["bitrate"] = hachoir_metadata.get("bit_rate")

        if hachoir_metadata.has("compression"):
            metadata["codec"] = str(
                hachoir_metadata.get("compression")
            )

    except Exception as error:
        metadata["error"] = str(error)

    return metadata

def extract_evidence_metadata(file_path: str) -> dict:
    extension = get_file_extension(file_path)

    if extension in IMAGE_EXTENSIONS:
        return extract_image_metadata(file_path)

    if extension in AUDIO_EXTENSIONS:
        return extract_audio_metadata(file_path)

    if extension in VIDEO_EXTENSIONS:
        return extract_video_metadata(file_path)

    return {}

def extract_text_from_file(file_path: str) -> str:
    extension = get_file_extension(file_path)

    if extension in [".txt", ".log", ".json", ".xml", ".ini", ".cfg", ".conf", ".md"]:
        with open(
            file_path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as file:
            return file.read()

    if extension == ".pdf":
        try:
            reader = PdfReader(file_path)
            text = ""
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
            return text
        except Exception as e:
            return f"[Error reading PDF: {e}]"

    if extension == ".csv":
        try:
            dataframe = pd.read_csv(file_path)
            return dataframe.to_string(index=False)
        except Exception:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as file:
                return file.read()

    if extension in [".docx", ".doc"]:
        try:
            import docx
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    paragraphs.append(" | ".join(row_data))
            return "\n".join(paragraphs)
        except Exception as e:
            return f"[Error reading DOCX: {e}]"

    return ""

def extract_text_lines(file_path: str) -> list[dict]:
    full_text = extract_text_from_file(file_path)
    if not full_text:
        return []
    lines = full_text.splitlines()
    return [{"line": idx + 1, "text": line} for idx, line in enumerate(lines)]

def verify_file_hash(file_path: str, stored_hash: str) -> bool:
    current_hash = calculate_file_hash(file_path)

    return current_hash == stored_hash


# Example usage / test
if __name__ == "__main__":
    # Replace with your actual file path and stored hash
    test_file = "example_evidence.mp4"
    stored_hash = "YOUR_STORED_SHA256_HASH_HERE"

    is_verified = verify_file_hash(test_file, stored_hash)

    if is_verified:
        print("Integrity: Verified ✅")
    else:
        print("Integrity: Failed ❌")
